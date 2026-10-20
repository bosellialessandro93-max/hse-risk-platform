from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
import re
p=Path(__file__).resolve().parent/'artifacts'
counts={}
for ext in ['pdf','docx','xlsx','pptx']:
 files=list(p.glob('*.'+ext));assert files,ext
 counts[ext]=len(files)
 for f in files:
  if ext=='pdf':
   reader=PdfReader(f);assert reader.pages
   if 'management' in f.name:assert len(reader.pages)>=8
   if 'A3' in f.name:assert abs(float(reader.pages[0].mediabox.width)-1191)<1
   if 'A4' in f.name:assert abs(float(reader.pages[0].mediabox.width)-842)<1
  if ext=='docx':assert Document(f).paragraphs
  if ext=='xlsx':assert len(load_workbook(f).sheetnames)>=2
  if ext=='pptx':
   with ZipFile(f) as z:
    slides=[n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$',n)];assert len(slides)>=8
    for n in slides:assert ET.fromstring(z.read(n)) is not None
ics=(p/'calendar.ics').read_text();assert ics.count('BEGIN:VEVENT')==100
photo_pdfs=list(p.glob('photo-*.pdf'));assert photo_pdfs
assert any(len(page.images)>0 for file in photo_pdfs for page in PdfReader(file).pages),'Evidenze fotografiche mancanti'
print('PASS export verificati:',counts,'ICS 100 eventi; PDF con evidenze fotografiche')

import json,math
saved=json.loads((p/'risk-saved.json').read_text())
for f in p.glob('risk-*.xlsx'):
 wb=load_workbook(f);assert math.isclose(wb['Sintesi management']['B5'].value,saved['result']['base'],abs_tol=1e-9)
 assert {'Contributi','Metodo','Storico','Action Plan','Scenari','Input completo'}<=set(wb.sheetnames)
 assert wb['Sintesi management']._images
for f in p.glob('risk-*.pptx'):
 with ZipFile(f) as z:
  assert any(n.startswith('ppt/charts/chart') and n.endswith('.xml') for n in z.namelist())
  slides=[n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$',n)]
  matching=f.with_name(f.stem+'-management.pdf');assert len(PdfReader(matching).pages)==len(slides)
for f in p.glob('sinottico-*.pdf'):
 reader=PdfReader(f);assert len(reader.pages)>1
 for page in reader.pages:
  assert abs(float(page.mediabox.width)-1191)<1
  assert 'Riferimento POS' in page.extract_text()
 text='\n'.join(page.extract_text() for page in reader.pages)
 assert 'DEMO lunga attività 0' in text and 'DEMO lunga attività 99' in text
for f in p.glob('sinottico-*.docx'):
 doc=Document(f);assert len(doc.tables[0].rows)==101
print('PASS numeri Excel riconciliati; grafico PPTX nativo; PPTX/PDF pagine coerenti; A3/Word 100 attività con intestazioni ripetute')

meeting_pdfs=list(p.glob('meeting-photo-*.pdf'));assert meeting_pdfs
for f in meeting_pdfs:assert sum(len(page.images) for page in PdfReader(f).pages)>=2
for f in p.glob('meeting-photo-*.docx'):
 with ZipFile(f) as z:assert len([n for n in z.namelist() if n.startswith('word/media/')])>=2
print('PASS foto delle osservazioni incluse nei verbali PDF e Word')
